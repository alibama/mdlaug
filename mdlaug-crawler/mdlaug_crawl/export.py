"""Export a run for sharing: an .xlsx workbook (Arial; summary uses formulas)."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import analysis, scoring

HDR = Font(name="Arial", bold=True, color="FFFFFF")
BODY = Font(name="Arial")
FILL = PatternFill("solid", fgColor="12332E")


def _sheet(ws, df):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font, c.fill = HDR, FILL
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(10, len(str(col)) + 2))
    ws.freeze_panes = "A2"


def export_xlsx(store, run_id, path, review_conf=0.7):
    t = analysis.site_situation_table(store, run_id, review_conf)
    wb = Workbook()
    ws = wb.active; ws.title = "Scores"
    codes = [c for c, _ in scoring.SITUATIONS]
    if not t.empty:
        wide = t.pivot_table(index=["institution", "platform", "library_type", "collection_url"],
                             columns="code", values="score", aggfunc="first").reindex(columns=codes).reset_index()
    else:
        wide = pd.DataFrame(columns=["institution", "platform", "library_type", "collection_url"] + codes)
    _sheet(ws, wide)
    n = len(wide)
    # Summary row with formulas (averages per situation)
    ws.append([])
    r = n + 3
    ws.cell(row=r, column=1, value="Average (1–7)").font = Font(name="Arial", bold=True)
    for j, _ in enumerate(codes, start=5):
        col = get_column_letter(j)
        ws.cell(row=r, column=j, value=f'=IFERROR(AVERAGE({col}2:{col}{n + 1}),"")').font = Font(name="Arial", bold=True)
    ws.cell(row=r + 1, column=1, value="Scores: reviewed final score where available, otherwise the pipeline's "
                                       "suggestion. See 'Detail' for status and confidence.").font = Font(name="Arial", italic=True)
    ws2 = wb.create_sheet("Detail")
    cols = ["institution", "platform", "library_type", "collection_url", "code", "level", "suggested_score",
            "final_score", "status", "confidence", "evidence_count", "methods"]
    d = t[cols].copy() if not t.empty else pd.DataFrame(columns=cols)
    if not d.empty:
        d["methods"] = d["methods"].apply(lambda m: ", ".join(m) if isinstance(m, list) else m)
    _sheet(ws2, d)
    ws3 = wb.create_sheet("Legend")
    for line in [
        "mDLAUG crawl export",
        f"Run: {run_id}",
        "suggested_score: pipeline suggestion on the 1–7 rubric (engine + Playwright probes + LLM judgments, confidence-weighted).",
        "final_score: the score a reviewer confirmed or set. Blank = not yet reviewed.",
        "status: automated (high-confidence, no review flag) · needs review · reviewed.",
        "Every suggested score is a starting point; situations marked 'needs review' require a person.",
    ]:
        ws3.append([line])
    for row in ws3.iter_rows():
        for c in row:
            c.font = BODY
    ws3.column_dimensions["A"].width = 120
    wb.save(path)
    return path
